# app.py — LA COCINA (toda la lógica vive aquí, NUNCA en el HTML)
# Academia Horizonte: lee los Excel de Insumos/, cruza, calcula y decide.
from pathlib import Path  # LADRILLO: rutas ancladas al propio archivo (tipo Path)
from urllib.parse import unquote  # LADRILLO: decodificar %20 y %C3%A9 que llegan en la URL
import hmac  # LADRILLO: comparar contraseñas sin filtrar tiempos
import io  # LADRILLO: memoria temporal para Excel, ZIP y PDFs
import os  # LADRILLO: leer variables de entorno (correos y contraseñas)
import re  # LADRILLO: limpiar nombres de archivo con expresiones regulares
import zipfile  # LADRILLO: empaquetar los PDF en un solo archivo
from datetime import date  # LADRILLO: fecha de expedición del certificado
from functools import wraps  # LADRILLO: conservar el nombre de las rutas decoradas
import openpyxl  # LADRILLO: librería para leer y escribir Excel .xlsx
from dotenv import load_dotenv  # LADRILLO: leer el archivo .env en la máquina local
from fpdf import FPDF  # LADRILLO: librería para crear archivos PDF
from openpyxl.styles import Font, PatternFill  # LADRILLO: estilos del Excel de salida
from openpyxl.utils import get_column_letter  # LADRILLO: letra de cada columna (A, B, C...)
# LADRILLO: servidor web, páginas, descargas y sesiones
from flask import (Flask, redirect, render_template, request, send_file,
                   session, url_for)

# LADRILLO: variables (cajas con nombre). Tipo Path y texto (str).
# Path(__file__) = ubicación de ESTE archivo, funciona sin importar desde dónde se arranque.
BASE = Path(__file__).resolve().parent

# LADRILLO: carga las variables del archivo .env (si existe) antes de leerlas.
# En Vercel no hay .env: las variables se configuran en el panel del proyecto.
load_dotenv(BASE / ".env")

RUTA_MAESTRO = BASE / "Insumos" / "Maestro_Estudiantes.xlsx"
RUTA_EVAL = BASE / "Insumos" / "Registro_Evaluaciones.xlsx"
DIR_CERTIFICADOS = BASE / "Certificados"  # LADRILLO: carpeta de salida (se crea sola)

# LADRILLO: límites de tamaño de las cargas.
# 4 MB por archivo y 4.5 MB por petición (Vercel corta por encima de 4.5 MB).
LIMITE_ARCHIVO = 4 * 1024 * 1024  # tipo int (bytes)
LIMITE_PETICION = int(4.5 * 1024 * 1024)  # tipo int (bytes)

# LADRILLO: texto fijo que describe de dónde salen los datos de la pantalla.
ORIGEN_FABRICA = ("Datos de fábrica (Insumos/Maestro_Estudiantes.xlsx "
                  "y Insumos/Registro_Evaluaciones.xlsx)")

# LADRILLO: constantes de color y textos fijos (tipo tupla / lista de textos)
COLOR_MARINO = (10, 42, 90)
COLOR_DORADO = (201, 162, 39)
COLOR_GRIS = (95, 95, 95)
COLOR_ROJO = (160, 40, 40)
COLOR_CREMA = (251, 246, 227)
MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio",
         "julio", "agosto", "setiembre", "octubre", "noviembre", "diciembre"]

app = Flask(__name__)  # LADRILLO: variable que guarda la aplicación web

# LADRILLO: clave para firmar la cookie de sesión (en producción va en el entorno)
app.secret_key = os.environ.get("SECRET_KEY", "clave-solo-para-desarrollo-local")

# LADRILLO: Vercel rechaza las peticiones de más de 4.5 MB. Flask responde 413
# y nosotros mostramos un aviso amable en vez de una página en blanco.
app.config["MAX_CONTENT_LENGTH"] = LIMITE_PETICION

# LADRILLO: opciones de seguridad de la cookie de sesión
app.config["SESSION_COOKIE_HTTPONLY"] = True  # el JavaScript no puede leer la cookie
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"  # evita envíos desde otros sitios


# LADRILLO: función auxiliar que lee una variable de entorno o usa un valor por defecto.
def leer_variable(nombre, por_defecto):
    # Devuelve el texto de la variable; si no existe o está vacía, el valor por defecto.
    valor = os.environ.get(nombre)  # tipo texto o None
    if valor is None or valor.strip() == "":
        return por_defecto  # tipo texto (modo demostración)
    return valor.strip()  # tipo texto


# LADRILLO: credenciales leídas del entorno. NUNCA se escriben en el código.
# Los valores de la derecha son SOLO para demostración si falta el .env.
ADMIN_CORREO = leer_variable("ADMIN_CORREO", "admin@academia.cr")
ADMIN_CLAVE = leer_variable("ADMIN_CLAVE", "admin123")
USUARIO_CORREO = leer_variable("USUARIO_CORREO", "usuario@academia.cr")
USUARIO_CLAVE = leer_variable("USUARIO_CLAVE", "usuario123")

# LADRILLO: hay modo demostración si falta al menos una de las cuatro variables.
MODO_DEMO = any(os.environ.get(n) in (None, "") for n in
                ("ADMIN_CORREO", "ADMIN_CLAVE", "USUARIO_CORREO", "USUARIO_CLAVE"))

# LADRILLO: lista de usuarios permitidos (tipo lista de dict). Sin base de datos.
USUARIOS = [
    {"correo": ADMIN_CORREO, "clave": ADMIN_CLAVE, "rol": "admin"},
    {"correo": USUARIO_CORREO, "clave": USUARIO_CLAVE, "rol": "usuario"},
]


# LADRILLO: función que compara correo y contraseña contra la lista de usuarios.
def validar_credenciales(correo, clave):
    # LADRILLO: bucle + condicional (revisa los dos usuarios configurados)
    for u in USUARIOS:  # u es tipo dict
        if u["correo"].lower() == correo.lower() and hmac.compare_digest(u["clave"], clave):
            return u
    return None  # si ninguno coincide, no hay usuario


# LADRILLO: decoradora (envuelve una ruta para exigir sesión antes de entrar).
def requiere_login(vista):
    @wraps(vista)  # conserva el nombre original de la función
    def envoltura(*args, **kwargs):
        # LADRILLO: condicional (si no hay correo en la sesión, no ha entrado)
        if "correo" not in session:
            return redirect(url_for("login"))
        return vista(*args, **kwargs)
    return envoltura


# LADRILLO: decoradora (envuelve una ruta para exigir el rol "admin").
def solo_admin(vista):
    @wraps(vista)  # conserva el nombre original de la función
    def envoltura(*args, **kwargs):
        # LADRILLO: condicional (cualquier rol distinto de "admin" recibe un 403)
        if session.get("rol") != "admin":
            return render_template("sin_permiso.html"), 403
        return vista(*args, **kwargs)
    return envoltura


# LADRILLO: función (bloque reutilizable: recibe datos, devuelve un resultado)
def calcular_promedio(valores):
    # LADRILLO: variables y tipos. float = número con decimales, int = entero.
    total = 0.0  # tipo float
    cantidad = 0  # tipo int
    # LADRILLO: bucle (repetir una acción por cada elemento de la lista)
    for v in valores:  # v es tipo float
        total = total + float(v)
        cantidad = cantidad + 1
    # LADRILLO: condicional (pregunta sí/no para decidir)
    if cantidad == 0:
        return 0.0
    return total / cantidad


# LADRILLO: función que decide el certificado según AGENTS.md (límites inclusivos)
def decidir_certificado(promedio, asistencia):
    # LADRILLO: condicionales encadenados (se revisan en orden)
    if asistencia < 80:  # tipo: comparación de float; menos de 80 = sin certificado
        return "Sin certificado"
    if promedio >= 70:  # tipo: comparación; 70 SÍ aprueba
        return "Aprobación"
    return "Participación"  # si no fue Sin ni Aprobación, es Participación


def cargar_datos(maestro=None, evaluaciones=None):
    # Lee los datos y deja todo listo para el Salón.
    # LADRILLO: los dos parámetros son OPCIONALES. Si llega un archivo subido (BytesIO),
    # se usa ese; si no llega nada, se lee el Excel de la carpeta Insumos.
    # Así la lógica de negocio se escribe UNA sola vez.
    fuente_maestro = maestro if maestro is not None else RUTA_MAESTRO
    fuente_evaluaciones = evaluaciones if evaluaciones is not None else RUTA_EVAL
    # LADRILLO: variables tipo lista (varios valores) y dict (tabla clave -> valor).
    wb_m = openpyxl.load_workbook(fuente_maestro, data_only=True)
    ws_m = wb_m.active
    maestro = {}  # tipo dict: clave (id, programa) -> datos del estudiante
    # LADRILLO: bucle por filas (saltamos el encabezado con [1:])
    for fila in list(ws_m.iter_rows(values_only=True))[1:]:
        if fila[0] is None:
            continue
        ident = str(fila[0]).strip()  # tipo texto: la ID siempre como texto
        maestro[(ident, str(fila[3]).strip())] = {
            "id": ident,  # tipo texto
            "nombre": str(fila[1]).strip(),  # tipo texto
            "correo": str(fila[2]).strip(),  # tipo texto
            "programa": str(fila[3]).strip(),  # tipo texto
            "cohorte": str(fila[4]).strip(),  # tipo texto
        }
    wb_m.close()

    wb_e = openpyxl.load_workbook(fuente_evaluaciones, data_only=True)
    ws_e = wb_e.active
    notas_por_clave = {}  # tipo dict: clave -> lista de notas (float)
    asis_por_clave = {}  # tipo dict: clave -> lista de asistencias (float)
    # LADRILLO: bucle para agrupar por Identificacion + Programa
    for fila in list(ws_e.iter_rows(values_only=True))[1:]:
        if fila[0] is None:
            continue
        clave = (str(fila[0]).strip(), str(fila[1]).strip())  # tipo tupla de textos
        # LADRILLO: condicional (si la clave es nueva, creo sus listas)
        if clave not in notas_por_clave:
            notas_por_clave[clave] = []
            asis_por_clave[clave] = []
        notas_por_clave[clave].append(float(fila[3]))  # tipo float
        asis_por_clave[clave].append(float(fila[4]))  # tipo float
    wb_e.close()

    # LADRILLO: bucle para la tabla principal (cruce interno: solo maestro CON evaluación)
    resultados = []  # tipo lista de dict
    for clave in notas_por_clave:
        # LADRILLO: condicional (si no está en el maestro, es huérfano: va a inconsistencias, no a la tabla)
        if clave not in maestro:
            continue
        promedio = calcular_promedio(notas_por_clave[clave])  # LADRILLO: llamada a función
        asistencia = calcular_promedio(asis_por_clave[clave])  # reutilizamos el promedio para asistencia
        tipo = decidir_certificado(promedio, asistencia)  # LADRILLO: llamada a función
        dato = maestro[clave]
        # LADRILLO: condicional para el color de la fila (verde / azul / rojo)
        if tipo == "Aprobación":
            clase = "aprobacion"
        elif tipo == "Participación":
            clase = "participacion"
        else:
            clase = "sin"
        resultados.append({
            "id": dato["id"],
            "nombre": dato["nombre"],
            "programa": dato["programa"],
            "cohorte": dato["cohorte"],
            "modulos": len(notas_por_clave[clave]),  # tipo int
            "promedio": round(promedio, 2),  # tipo float
            "asistencia": round(asistencia, 2),  # tipo float
            "tipo": tipo,  # tipo texto
            "clase": clase,  # tipo texto (para el color)
        })
    # Orden por programa y nombre (sin lógica en el HTML)
    resultados.sort(key=lambda r: (r["programa"], r["nombre"]))

    # Contadores para el encabezado (se calculan aquí, el HTML solo los muestra)
    contadores = {"Aprobación": 0, "Participación": 0, "Sin certificado": 0}  # tipo dict
    # LADRILLO: bucle + condicionales para contar
    for r in resultados:
        if r["tipo"] == "Aprobación":
            contadores["Aprobación"] = contadores["Aprobación"] + 1
        elif r["tipo"] == "Participación":
            contadores["Participación"] = contadores["Participación"] + 1
        else:
            contadores["Sin certificado"] = contadores["Sin certificado"] + 1

    programas = sorted({m["programa"] for m in maestro.values()})  # tipo lista de textos

    # Las 3 inconsistencias detectadas en la validación
    prom_h = round(calcular_promedio(notas_por_clave.get(("999880777", "Técnico en IA Aplicada"), [0])), 2)
    asis_h = round(calcular_promedio(asis_por_clave.get(("999880777", "Técnico en IA Aplicada"), [0])), 2)
    inconsistencias = [
        {"titulo": "Maestro SIN evaluación (1)",
         "detalle": "Espinoza Leon Javier | 304560321 | Técnico en IA Aplicada: está en el maestro pero no tiene ninguna fila de notas."},
        {"titulo": "Evaluación SIN maestro (1)",
         "detalle": f"ID 999880777 | Técnico en IA Aplicada: tiene 4 módulos (prom {prom_h}, asis {asis_h}) pero no existe en el maestro, no se puede certificar."},
        {"titulo": "Posible abandono con ceros (1)",
         "detalle": "Zuñiga Arias Fernanda | 111230984: notas 88, 90, 0, 0 y asistencias 70, 55, 20, 0 (prom 44.5, asis 36.25). Los ceros cuentan en el promedio."},
    ]
    return resultados, contadores, programas, inconsistencias


# ---------------------------------------------------------------
# CARGA DE ARCHIVOS Y EXPORTACIÓN A EXCEL (todo en memoria)
# ---------------------------------------------------------------
def leer_archivo_subido(archivo, etiqueta):
    # LADRILLO: valida un archivo del formulario y devuelve sus bytes.
    # Devuelve tres valores: contenido (bytes o None), nombre visible y mensaje de error.
    if archivo is None or not archivo.filename:
        return None, "", f"Falta el archivo de {etiqueta}."
    # LADRILLO: nos quedamos solo con el nombre, sin la ruta que a veces manda el navegador.
    nombre = archivo.filename.replace("\\", "/").split("/")[-1].strip()
    if not nombre.lower().endswith(".xlsx"):
        return None, nombre, f"El archivo de {etiqueta} debe ser un Excel con extensión .xlsx."
    # LADRILLO: leemos como máximo 4 MB + 1 byte para saber si se pasó del límite.
    contenido = archivo.read(LIMITE_ARCHIVO + 1)  # tipo bytes
    if len(contenido) > LIMITE_ARCHIVO:
        return None, nombre, (f"El archivo de {etiqueta} pesa más de 4 MB. "
                              "Súbelo de nuevo con un tamaño menor.")
    if len(contenido) == 0:
        return None, nombre, f"El archivo de {etiqueta} está vacío."
    return contenido, nombre, None


def construir_excel_resultados(resultados):
    # LADRILLO: arma el Excel de resultados en memoria, sin escribir en el disco.
    libro = openpyxl.Workbook()
    hoja = libro.active
    hoja.title = "Resultados"

    encabezados = ["Identificación", "Nombre completo", "Programa", "Cohorte",
                   "Módulos", "Promedio", "Asistencia (%)", "Certificado"]
    hoja.append(encabezados)

    # LADRILLO: bucle para pasar cada resultado a una fila del Excel.
    for r in resultados:
        hoja.append([r["id"], r["nombre"], r["programa"], r["cohorte"],
                     r["modulos"], r["promedio"], r["asistencia"], r["tipo"]])

    # Estilo de la fila de títulos: fondo marino y letra blanca.
    for celda in hoja[1]:
        celda.fill = PatternFill("solid", fgColor="0A2A5A")
        celda.font = Font(bold=True, color="FFFFFF")

    # Anchos cómodos para leer el archivo.
    anchos = [16, 32, 34, 12, 10, 10, 14, 20]
    for numero, ancho in enumerate(anchos, start=1):
        hoja.column_dimensions[get_column_letter(numero)].width = ancho

    # Dos decimales en promedio y asistencia.
    for fila in hoja.iter_rows(min_row=2, min_col=6, max_col=7):
        for celda in fila:
            celda.number_format = "0.00"
    hoja.freeze_panes = "A2"

    # LADRILLO: ¿Qué es io.BytesIO? Es como una caja portátil que vive en la memoria RAM:
    # guardamos el archivo "dentro de la caja" en vez de en una carpeta del disco, y openpyxl
    # lo lee o lo escribe igual que si fuera un archivo real. Al terminar, la caja se tira.
    buffer = io.BytesIO()  # tipo BytesIO (archivo imaginario en memoria)
    libro.save(buffer)
    buffer.seek(0)  # rebobinar para que la descarga empiece en la primera fila
    return buffer


def enviar_excel_resultados(resultados, nombre_archivo):
    # LADRILLO: send_file envía la caja de memoria como descarga, sin tocar el disco.
    buffer = construir_excel_resultados(resultados)
    return send_file(
        buffer,
        as_attachment=True,
        download_name=nombre_archivo,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


# ---------------------------------------------------------------
# GENERACIÓN DE PDF (un archivo por estudiante con evaluación)
# ---------------------------------------------------------------
def ajustar_fuente(pdf, texto, fuente, estilo, tam_max, tam_min, ancho):
    # LADRILLO: función que reduce el tamaño de letra hasta que el texto quepa.
    tam = tam_max  # tipo int
    pdf.set_font(fuente, estilo, tam)
    # LADRILLO: bucle + condicional (mientras sea muy ancho y se pueda reducir)
    while pdf.get_string_width(texto) > ancho and tam > tam_min:
        tam = tam - 1
        pdf.set_font(fuente, estilo, tam)
    return tam  # tipo int


def nombre_archivo(r):
    # LADRILLO: función que arma un nombre de archivo válido para Windows.
    texto = f"{r['id']} - {r['nombre']} - {r['programa']}.pdf"  # tipo texto
    # Quita los caracteres prohibidos por Windows: < > : " / \ | ? *
    return re.sub(r'[<>:"/\\|?*]', "", texto)


def texto_centrado(pdf, y, texto, fuente, estilo, tam, color, ancho=273, alto=8):
    # LADRILLO: función auxiliar para escribir una línea centrada.
    pdf.set_font(fuente, estilo, tam)
    pdf.set_text_color(*color)
    pdf.set_xy(12, y)
    pdf.cell(ancho, alto, texto, align="C")


def construir_certificado_pdf(r):
    # LADRILLO: función que devuelve los bytes del PDF de UN estudiante.
    # A4 horizontal: 297 mm de ancho x 210 mm de alto.
    pdf = FPDF(orientation="L", unit="mm", format="A4")
    pdf.set_auto_page_break(False)
    pdf.set_margins(0, 0, 0)
    pdf.add_page()

    # Fondo blanco y doble borde (marino por fuera, dorado por dentro)
    pdf.set_fill_color(255, 255, 255)
    pdf.rect(0, 0, 297, 210, "F")
    pdf.set_draw_color(*COLOR_MARINO)
    pdf.set_line_width(1.4)
    pdf.rect(8, 8, 281, 194)
    pdf.set_draw_color(*COLOR_DORADO)
    pdf.set_line_width(0.6)
    pdf.rect(12, 12, 273, 186)

    # Esquinas decorativas doradas (cuatro escuadras)
    pdf.set_draw_color(*COLOR_DORADO)
    pdf.set_line_width(1.2)
    for x, y, dx, dy in [(18, 18, 12, 0), (18, 18, 0, 12),
                         (279, 18, -12, 0), (279, 18, 0, 12),
                         (18, 192, 12, 0), (18, 192, 0, -12),
                         (279, 192, -12, 0), (279, 192, 0, -12)]:
        pdf.line(x, y, x + dx, y + dy)

    # Encabezado: nombre de la academia y raya dorada
    texto_centrado(pdf, 22, "ACADEMIA HORIZONTE", "Times", "B", 22, COLOR_MARINO, alto=10)
    pdf.set_draw_color(*COLOR_DORADO)
    pdf.set_line_width(0.8)
    pdf.line(108, 35, 189, 35)

    # LADRILLO: condicional para elegir el texto según el tipo de certificado
    if r["tipo"] == "Aprobación":
        subtitulo = "CERTIFICADO DE APROBACIÓN"
        introduccion = "Se otorga el presente certificado a"
        color_titulo = COLOR_MARINO
    elif r["tipo"] == "Participación":
        subtitulo = "CONSTANCIA DE PARTICIPACIÓN"
        introduccion = "Se otorga la presente constancia a"
        color_titulo = COLOR_MARINO
    else:
        subtitulo = "REGISTRO ACADÉMICO - SIN CERTIFICADO"
        introduccion = "Se registra el desempeño académico de"
        color_titulo = COLOR_ROJO

    texto_centrado(pdf, 39, subtitulo, "Helvetica", "B", 12, color_titulo)
    texto_centrado(pdf, 53, introduccion, "Times", "I", 11, COLOR_GRIS)

    # Nombre del estudiante (se reduce si es muy largo)
    ajustar_fuente(pdf, r["nombre"], "Times", "B", 30, 18, 235)
    pdf.set_text_color(*COLOR_MARINO)
    pdf.set_xy(12, 63)
    pdf.cell(273, 16, r["nombre"], align="C")

    # Raya dorada bajo el nombre
    pdf.set_draw_color(*COLOR_DORADO)
    pdf.set_line_width(0.8)
    pdf.line(88, 82, 209, 82)

    texto_centrado(pdf, 85, "por haber cursado y completado el programa", "Times", "I", 11, COLOR_GRIS)

    # Programa y cohorte
    ajustar_fuente(pdf, r["programa"], "Helvetica", "B", 16, 11, 235)
    pdf.set_text_color(*COLOR_MARINO)
    pdf.set_xy(12, 94)
    pdf.cell(273, 9, r["programa"], align="C")
    texto_centrado(pdf, 104, f"Cohorte {r['cohorte']}", "Helvetica", "", 11, COLOR_GRIS)

    # Recuadro crema con las tres métricas
    pdf.set_fill_color(*COLOR_CREMA)
    pdf.set_draw_color(*COLOR_DORADO)
    pdf.set_line_width(0.4)
    pdf.rect(52, 118, 193, 28, "DF")

    metricas = [
        ("PROMEDIO", f"{r['promedio']:.2f}"),
        ("ASISTENCIA", f"{r['asistencia']:.2f}%"),
        ("MÓDULOS", str(r["modulos"])),
    ]
    # LADRILLO: bucle para dibujar las tres columnas del recuadro
    for i, (etiqueta, valor) in enumerate(metricas):
        centro_x = 52 + 193 * (i + 0.5) / 3  # tipo float
        pdf.set_font("Helvetica", "B", 8)
        pdf.set_text_color(*COLOR_GRIS)
        pdf.set_xy(centro_x - 30, 122)
        pdf.cell(60, 5, etiqueta, align="C")
        pdf.set_font("Helvetica", "B", 15)
        pdf.set_text_color(*COLOR_MARINO)
        pdf.set_xy(centro_x - 30, 128)
        pdf.cell(60, 9, valor, align="C")
        # Separadores verticales entre columnas
        if i < 2:
            pdf.set_draw_color(*COLOR_DORADO)
            pdf.line(52 + 193 * (i + 1) / 3, 121, 52 + 193 * (i + 1) / 3, 143)

    # Nota extra cuando no hay certificado
    if r["tipo"] == "Sin certificado":
        texto_centrado(pdf, 150, "No cumple los requisitos de certificación (promedio >= 70 y asistencia >= 80).",
                       "Helvetica", "B", 9, COLOR_ROJO)

    # Identificación a la izquierda y fecha a la derecha
    fecha = date.today()
    fecha_texto = f"{fecha.day} de {MESES[fecha.month - 1]} de {fecha.year}"
    pdf.set_font("Helvetica", "", 10)
    pdf.set_text_color(*COLOR_GRIS)
    pdf.set_xy(28, 160)
    pdf.cell(90, 6, f"Identificación: {r['id']}", align="L")
    pdf.set_xy(179, 160)
    pdf.cell(90, 6, f"Expedido el {fecha_texto}", align="R")

    # Firmas
    pdf.set_draw_color(*COLOR_MARINO)
    pdf.set_line_width(0.5)
    pdf.line(55, 180, 125, 180)
    pdf.line(172, 180, 242, 180)
    texto_centrado(pdf, 182, "Coordinación Académica", "Helvetica", "", 9, COLOR_GRIS, ancho=70)
    pdf.set_xy(55, 182)
    pdf.cell(70, 5, "Coordinación Académica", align="C")
    pdf.set_xy(172, 182)
    pdf.cell(70, 5, "Dirección Académica", align="C")
    texto_centrado(pdf, 190, "Academia Horizonte - Formación técnica para empresas", "Times", "I", 8, COLOR_GRIS)

    # fpdf2 devuelve los bytes si no se indica archivo de salida
    return bytes(pdf.output())


def generar_certificado_pdf(r, ruta):
    # LADRILLO: función que guarda el PDF de un estudiante en disco.
    contenido = construir_certificado_pdf(r)  # tipo bytes
    ruta.write_bytes(contenido)
    return ruta


def generar_todos_los_pdf():
    # LADRILLO: función que genera un PDF por cada estudiante (los 23 con evaluación).
    resultados, _, _, _ = cargar_datos()
    DIR_CERTIFICADOS.mkdir(exist_ok=True)  # crea Certificados/ si no existe
    rutas = []  # tipo lista de Path
    # LADRILLO: bucle sobre todos los resultados
    for r in resultados:
        ruta = DIR_CERTIFICADOS / nombre_archivo(r)
        generar_certificado_pdf(r, ruta)
        rutas.append(ruta)
    return rutas


# ---------------------------------------------------------------
# RUTAS WEB
# ---------------------------------------------------------------
def buscar_resultado(ident, programa):
    # LADRILLO: en Vercel el programa puede llegar codificado (%20, %C3%A9).
    # Lo decodificamos para compararlo con el texto limpio que viene del Excel.
    programa = unquote(programa)
    # LADRILLO: busca un resultado por su llave Identificación + Programa.
    resultados, _, _, _ = cargar_datos()
    for r in resultados:  # tipo dict
        if r["id"] == ident and r["programa"] == programa:
            return r
    return None


def mostrar_tablero(resultados, contadores, programas, inconsistencias,
                    programa_actual="Todos", origen_datos=ORIGEN_FABRICA,
                    datos_subidos=False, error=None):
    # LADRILLO: prepara el Salón. Si el rol no es admin, las inconsistencias no salen.
    # Si la tabla viene de archivos subidos, tampoco mostramos los avisos de fábrica.
    if session.get("rol") != "admin" or datos_subidos:
        inconsistencias = []
    # LADRILLO: condicional para el filtro por programa (llega como ?programa=...)
    if programa_actual == "Todos":
        filtrados = resultados
    else:
        filtrados = [r for r in resultados if r["programa"] == programa_actual]
    # LADRILLO: el HTML solo muestra; todos los datos llegan calculados desde aquí.
    return render_template("index.html", resultados=filtrados, contadores=contadores,
                           programas=programas, programa_actual=programa_actual,
                           inconsistencias=inconsistencias, origen_datos=origen_datos,
                           datos_subidos=datos_subidos, error=error)


# LADRILLO: deja disponibles el correo y el rol en todas las plantillas.
@app.context_processor
def datos_de_sesion():
    return {
        "correo_sesion": session.get("correo", ""),
        "rol_sesion": session.get("rol", ""),
    }


# LADRILLO: pantalla de entrada (GET muestra el formulario, POST valida).
@app.route("/login", methods=["GET", "POST"])
def login():
    # LADRILLO: condicional (si ya entró, no tiene sentido mostrar el login)
    if "correo" in session:
        return redirect(url_for("index"))

    error = None  # tipo texto o None
    correo_ingresado = ""  # tipo texto (para no borrar el correo si falla)
    # LADRILLO: condicional según el método del formulario
    if request.method == "POST":
        correo_ingresado = request.form.get("correo", "").strip()  # tipo texto
        clave = request.form.get("clave", "")  # tipo texto
        usuario = validar_credenciales(correo_ingresado, clave)  # tipo dict o None
        if usuario is not None:
            session.clear()  # borra cualquier dato anterior
            session["correo"] = usuario["correo"]  # guarda el correo en la sesión
            session["rol"] = usuario["rol"]  # guarda el rol en la sesión
            return redirect(url_for("index"))
        error = "Correo o contraseña incorrectos."

    # El Salón solo muestra: las credenciales de demo vienen preparadas desde aquí.
    return render_template("login.html", error=error, correo_ingresado=correo_ingresado,
                           modo_demo=MODO_DEMO, demo_admin=ADMIN_CORREO,
                           demo_usuario=USUARIO_CORREO,
                           demo_admin_clave=ADMIN_CLAVE,
                           demo_usuario_clave=USUARIO_CLAVE)


# LADRILLO: ruta para cerrar sesión (borra todo y vuelve a la entrada).
@app.route("/salir")
def salir():
    session.clear()  # elimina correo y rol de la cookie de sesión
    return redirect(url_for("login"))


# LADRILLO: ruta web (qué página se muestra al visitar una dirección)
@app.route("/")
@requiere_login
def index():
    programa_actual = request.args.get("programa", "Todos")  # tipo texto
    # LADRILLO: sin archivos subidos, cargar_datos() lee los Excel de Insumos.
    return mostrar_tablero(*cargar_datos(), programa_actual=programa_actual)


# LADRILLO: recibe los dos Excel, los procesa EN MEMORIA y muestra o descarga el resultado.
@app.route("/cargar", methods=["POST"])
@requiere_login
@solo_admin
def cargar():
    # LADRILLO: validamos y leemos el primer archivo (máximo 4 MB).
    contenido_maestro, nombre_maestro, error = leer_archivo_subido(
        request.files.get("maestro"), "maestro de estudiantes")
    # LADRILLO: si el primero salió bien, validamos también el segundo.
    if error is None:
        contenido_eval, nombre_eval, error = leer_archivo_subido(
            request.files.get("evaluaciones"), "registro de evaluaciones")

    # LADRILLO: si falta un archivo, no es Excel o pesa demasiado, avisamos en rojo.
    if error is not None:
        return mostrar_tablero(*cargar_datos(), error=error)

    # LADRILLO: io.BytesIO es la "caja" en memoria donde openpyxl lee el Excel subido.
    # No se escribe nada en el disco (en Vercel el disco es de solo lectura).
    try:
        resultados, contadores, programas, inconsistencias = cargar_datos(
            io.BytesIO(contenido_maestro), io.BytesIO(contenido_eval))
    except Exception as detalle:
        # LADRILLO: si el Excel está dañado o no tiene las columnas esperadas, no rompemos la página.
        app.logger.warning("No se pudieron leer los Excel subidos: %s", detalle)
        return mostrar_tablero(
            *cargar_datos(),
            error=("No pudimos leer los Excel subidos. Revisa que sean los archivos "
                   "correctos y que tengan las columnas esperadas."))

    # LADRILLO: el botón pulsado llega en el campo "accion" del formulario.
    accion = request.form.get("accion", "ver")  # tipo texto
    if accion == "excel":
        # LADRILLO: descarga armada en memoria, sin guardar el archivo en el servidor.
        return enviar_excel_resultados(resultados, "resultados_subidos.xlsx")

    # LADRILLO: mostramos la tabla con el origen y el aviso de que los datos son temporales.
    origen = f"Archivos subidos: {nombre_maestro} y {nombre_eval}"
    return mostrar_tablero(resultados, contadores, programas, inconsistencias,
                           origen_datos=origen, datos_subidos=True)


# LADRILLO: descarga de los resultados calculados con los Excel de fábrica.
@app.route("/resultados-fabrica.xlsx")
@requiere_login
@solo_admin
def resultados_fabrica():
    resultados, _, _, _ = cargar_datos()
    return enviar_excel_resultados(resultados, "resultados_fabrica.xlsx")


# LADRILLO: si la petición completa supera 4.5 MB, Flask lanza 413.
# Lo atrapamos para mostrar un aviso rojo en el tablero y no una página en blanco.
@app.errorhandler(413)
def carga_demasiado_grande(_error):
    mensaje = ("La carga supera el límite de 4.5 MB por envío (Vercel no acepta más). "
               "Sube archivos más livianos, de máximo 4 MB cada uno.")
    # LADRILLO: sin sesión de admin no hay tablero que mostrar; mandamos a entrar.
    if session.get("rol") != "admin":
        return redirect(url_for("login"))
    return mostrar_tablero(*cargar_datos(), error=mensaje), 413


# LADRILLO: ruta exclusiva del administrador (protección real en el servidor).
@app.route("/inconsistencias")
@requiere_login
@solo_admin
def inconsistencias():
    _, _, _, lista = cargar_datos()
    return render_template("inconsistencias.html", inconsistencias=lista)


@app.route("/certificado/<ident>/<path:programa>/pdf")
@requiere_login
def certificado_pdf(ident, programa):
    # LADRILLO: busca al estudiante y arma SU PDF en memoria, sin escribir en disco.
    elegido = buscar_resultado(ident, programa)
    if elegido is None:
        return "Certificado no encontrado", 404
    contenido = construir_certificado_pdf(elegido)  # tipo bytes
    # LADRILLO: io.BytesIO permite enviar esos bytes como si fueran un archivo.
    return send_file(
        io.BytesIO(contenido),
        mimetype="application/pdf",
        as_attachment=True,
        download_name=nombre_archivo(elegido),
    )


@app.route("/certificado/<ident>/<path:programa>")
@requiere_login
def certificado(ident, programa):
    elegido = buscar_resultado(ident, programa)
    if elegido is None:
        return "Certificado no encontrado", 404
    return render_template("certificado.html", r=elegido)


if __name__ == "__main__":
    # Arranca el Salón en http://localhost:5000 ( codificación UTF-8 para tildes )
    app.run(host="127.0.0.1", port=5000, debug=False)
